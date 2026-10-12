"""Build the evidence-aligned project report as a Word document."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Revised_Project_Report_Chapters_1_to_5.docx"


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, rows):
    for row in rows:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(row)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    doc.add_paragraph()
    return table


def chapter_title(doc, number, title):
    doc.add_page_break()
    h = doc.add_heading(f"CHAPTER {number}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h = doc.add_heading(title.upper(), level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INTELLIGENT WASTE CLASSIFICATION SYSTEM")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Revised Chapters One to Five: Evidence-Aligned Project Report").italic = True
    doc.add_paragraph()
    add_para(doc, "Prepared from the project repository and the supplied Chapters 1–3 document. "
                  "This revision separates implemented functionality from experimental findings so that no unmeasured performance is presented as a result.")

    add_heading(doc, "Verified project status", 1)
    add_para(doc, "Repository inspection on 6 October 2026 found no trained .keras model, metrics JSON, training-history plot, confusion-matrix image, or non-placeholder dataset image. TensorFlow is also not installed in the current workspace. Consequently, the project implementation exists, but neither the baseline CNN nor the MobileNetV2 model has verifiable training or test-set results in this copy of the project.")
    add_table(doc, ["Item checked", "Verified status", "Implication for this report"], [
        ["data/train, data/val and data/test", "Only class folders and .gitkeep placeholder files", "No dataset is available for training or evaluation."],
        ["models/", "Directory contains no output artifacts", "No trained model or reported metric can be verified."],
        ["Training notebook", "All execution counts are null; no saved outputs", "Notebook is a training plan, not evidence of a completed run."],
        ["TensorFlow runtime", "Not installed", "Training/evaluation cannot be reproduced in this workspace."],
    ])

    chapter_title(doc, 1, "Introduction")
    add_heading(doc, "1.1 Background of the Study", 2)
    add_para(doc, "Solid-waste management depends on collection, separation, recovery and safe disposal. Source separation improves the quality of materials available for recycling and composting, while mixed waste makes recovery more difficult. The World Bank identifies waste prevention, source reduction, collection, recycling, composting and environmentally sound disposal as complementary components of effective municipal waste management (World Bank, 2025). In Nigeria, the need for better source segregation and treatment infrastructure makes a low-cost classification aid worth investigating.")
    add_para(doc, "Computer vision can support preliminary material recognition by mapping an image of an item to a predefined class. Convolutional neural networks (CNNs) learn image features directly from training examples, while transfer learning adapts representations learned on a large source dataset to a smaller target task (Pan & Yang, 2010). This study implements and compares a custom CNN with MobileNetV2, a mobile-oriented architecture built around inverted residual blocks and depthwise convolutions (Sandler et al., 2018).")
    add_heading(doc, "1.2 Statement of the Problem", 2)
    add_para(doc, "Manual or absent segregation can expose handlers to hazards and reduce material recovery. However, an image classifier trained on a restricted or curated dataset cannot be assumed to work under varied real-world conditions. The central problem is therefore to develop and evaluate—using a documented, held-out test set—a prototype that distinguishes plastic, paper, metal, glass and organic waste, while also considering the computational cost of deployment. At the time of this revision, the required data and measured model results are not yet available; the claims of performance remain research questions rather than findings.")
    add_heading(doc, "1.3 Aim and Objectives", 2)
    add_para(doc, "The aim is to develop and empirically evaluate an image-based waste-classification prototype suitable for initial investigation in a Nigerian waste-management context.")
    add_bullets(doc, [
        "Assemble a consented and documented five-class image dataset, including locally captured images where available.",
        "Preprocess and augment only the training images, while preserving separate validation and test sets.",
        "Train a baseline CNN and a MobileNetV2 transfer-learning model under the same split and class mapping.",
        "Evaluate both models on the held-out test set using accuracy, macro precision, macro recall, macro F1-score, a class-level report, a confusion matrix and average inference time.",
        "Integrate the selected trained model into a web interface and test the upload-to-prediction flow."
    ])
    add_heading(doc, "1.4 Research Questions", 2)
    add_bullets(doc, [
        "How accurately do the two models classify the five specified waste categories on a held-out test set?",
        "How do macro precision, macro recall, macro F1-score and the confusion matrix qualify the overall accuracy?",
        "Which of the two models offers the better accuracy–inference-time trade-off for the observed deployment environment?",
        "Does the web application return a valid class and probability distribution for supported image uploads?"
    ])
    add_heading(doc, "1.5 Significance, Scope and Delimitation", 2)
    add_para(doc, "The contribution is a reproducible prototype and a transparent comparison, not a claim that the system is ready to control a waste-sorting line. The scope is single-image classification into plastic, paper, metal, glass and organic classes. It excludes multi-object detection, hazardous-waste identification, automated mechanical sorting, and a claim of generalisation beyond the documented test data. Predictions should be treated as decision support and checked by a human where an error could create safety, environmental or financial harm.")
    add_heading(doc, "1.6 Review of the Supplied Chapter One", 2)
    add_para(doc, "The original chapter establishes a useful practical motivation, but its objectives and research questions named CNN, ResNet and MobileNet even though the repository implements only a custom CNN and MobileNetV2. The revised version removes the unimplemented ResNet comparison and changes future-tense performance assertions into measurable evaluation questions. It also avoids saying that the model is already locally trained, because no such artifact was supplied.")

    chapter_title(doc, 2, "Literature Review")
    add_heading(doc, "2.1 Waste Classification as an Image-Recognition Task", 2)
    add_para(doc, "Benchmark datasets have enabled rapid experimentation in waste recognition. TrashNet, for example, was developed as a dataset for classifying recyclable material images (Thung & Yang, 2016). Such benchmarks are useful for controlled comparison, but their labels, backgrounds and class distribution may differ from images collected in the target setting. A credible study must therefore record data provenance, class counts, split procedure, duplicate checks and the conditions represented in each split.")
    add_heading(doc, "2.2 CNNs and Transfer Learning", 2)
    add_para(doc, "CNNs progressively learn local and higher-level visual patterns from image pixels. A small custom CNN is a meaningful baseline because it reveals the performance obtained without pretraining. Transfer learning is often attractive for limited labelled data because a model pretrained on a broad source task can be adapted to the target task (Pan & Yang, 2010). MobileNetV2 is appropriate for comparison because it explicitly targets a balance between predictive performance, parameters and latency (Sandler et al., 2018). The architecture choice does not itself establish superiority; it must be assessed with the same data split and metrics.")
    add_heading(doc, "2.3 Evaluation and Reporting", 2)
    add_para(doc, "Accuracy alone can conceal poor recognition of a minority class. Macro-averaged precision, recall and F1-score give each class equal weight, making them suitable complements where class counts differ. A confusion matrix identifies the pairs of classes that are confused. The scikit-learn toolkit provides standard implementations of these measures (Pedregosa et al., 2011). Inference time should be reported with the hardware, software version, batch size and timing method, because latency cannot be meaningfully compared without those conditions.")
    add_heading(doc, "2.4 Gap and Position of the Present Study", 2)
    add_para(doc, "The practical gap addressed here is not merely the absence of another classifier; it is the need for a transparent, reproducible comparison that connects a documented dataset, a held-out evaluation and a usable interface. The study will only claim local relevance to the extent that locally captured, ethically obtained images are actually included and separately described. It will not infer field performance from a public benchmark alone.")
    add_heading(doc, "2.5 Review of the Supplied Chapter Two", 2)
    add_para(doc, "The supplied literature review is well structured and identifies important themes—dataset shift, efficiency and deployment. It should nevertheless be revised before submission: several assertions report precise accuracy figures without supplying the dataset split, evaluation protocol or independently verifiable primary source; some references are incomplete; and the chapter promises explainability and class-balancing work that the present repository does not implement. The present revision retains only claims that are necessary to motivate the actual two-model study and cites foundational or institutional sources that can be checked.")

    chapter_title(doc, 3, "System Design and Methodology")
    add_heading(doc, "3.1 Research Design", 2)
    add_para(doc, "The study follows an iterative data-mining and prototype-development process: define the evaluation objective, acquire and inspect data, prepare class-consistent splits, train candidate models, evaluate only on untouched test data, integrate the selected model, and document limitations. This is consistent with the iterative CRISP-DM framework (Chapman et al., 2000).")
    add_heading(doc, "3.2 Dataset and Split Protocol", 2)
    add_para(doc, "Five folders—glass, metal, organic, paper and plastic—are expected in each of train, validation and test. The supplied preparation script is designed to shuffle each class using seed 42 and create an 80:10:10 split. Before training, the final report must record actual images per class and source, the number removed as unreadable or duplicate, consent/permission arrangements for local photographs, and any class-name mapping (for example, whether cardboard is included in paper). The repository currently has zero usable images, so the proposed counts in the original Chapter Three are targets rather than achieved data.")
    add_heading(doc, "3.3 Preprocessing and Augmentation", 2)
    add_para(doc, "The implemented loaders resize images to 224 × 224 pixels and rescale pixel values to [0, 1]. Training images receive rotation up to 20 degrees, horizontal flipping, 15% zoom, brightness adjustment from 0.8 to 1.2, and horizontal/vertical shifts of 10%. Validation and test images are rescaled but not augmented. The same [0, 1] preprocessing is used in the Flask and Streamlit prediction paths. These details replace broader or inconsistent statements in the supplied chapter.")
    add_heading(doc, "3.4 Implemented Models", 2)
    add_table(doc, ["Model", "Implemented configuration", "Training plan in code"], [
        ["Baseline CNN", "Three Conv2D/ReLU/MaxPool blocks (32, 64 and 128 filters), Flatten, Dense(128), Dropout(0.4), softmax output.", "Adam optimizer; categorical cross-entropy; up to 50 epochs; early stopping; learning-rate reduction; best validation-accuracy checkpoint."],
        ["MobileNetV2", "ImageNet-pretrained base; global average pooling; Dense(128); Dropout(0.3); softmax output.", "Frozen base for up to 15 epochs at 1e-3, then layers from index 100 fine-tuned for up to 20 epochs at 1e-5; early stopping; best validation-accuracy checkpoint."]
    ])
    add_heading(doc, "3.5 Evaluation Procedure", 2)
    add_para(doc, "The evaluation script loads a saved .keras model, predicts the non-shuffled test generator and reports accuracy, macro precision, macro recall, macro F1-score, an sklearn per-class classification report, a confusion matrix and mean prediction time per image. It writes a metrics JSON file and a confusion-matrix PNG. The model selected for deployment must be chosen after this comparison, not assumed in advance. A future run should also record model file size, trainable parameter count, hardware and total training time.")
    add_heading(doc, "3.6 Review of the Supplied Chapter Three", 2)
    add_para(doc, "The original methodology is broadly compatible with the repository, but three corrections are required. First, it says both models run for a maximum of 50 epochs; MobileNetV2 actually has two phases of 15 and 20 epochs. Second, it mentions SavedModel or .h5, while the scripts save .keras files. Third, it describes data collection, local images and system testing as though completed, but the repository contains no evidence of these activities. They are retained as planned procedures until documented outputs are produced.")

    chapter_title(doc, 4, "Implementation, Verification and Results")
    add_heading(doc, "4.1 Introduction", 2)
    add_para(doc, "This chapter reports only evidence available in the supplied project folder. It distinguishes source-code implementation from model training and experimental evaluation. This distinction is essential: a training script or a web interface does not constitute a trained model or a measured result.")
    add_heading(doc, "4.2 Implementation Completed", 2)
    add_table(doc, ["Component", "Evidence found", "Status"], [
        ["Data preparation", "prepare_dataset.py validates images, converts copies to JPEG, and creates per-class 80:10:10 splits.", "Implemented; not executed on supplied data."],
        ["Baseline model", "src/train_baseline_cnn.py defines, trains and checkpoints a custom CNN.", "Implemented; no checkpoint exists."],
        ["Transfer-learning model", "src/train_mobilenet.py defines two-phase MobileNetV2 training.", "Implemented; no checkpoint exists."],
        ["Evaluation", "src/evaluate.py calculates aggregate and per-class metrics and saves a confusion matrix.", "Implemented; no metrics output exists."],
        ["Flask and Streamlit interfaces", "Both load a saved model, preprocess an upload and return the top class, confidence and all probabilities.", "Implemented; unavailable until a model file is supplied."],
        ["Automated tests", "Test files exist for data utilities and both web applications.", "Not executed in this workspace because required runtime dependencies are absent."],
    ])
    add_heading(doc, "4.3 Dataset and Training Verification", 2)
    add_para(doc, "The train, validation and test directories contain the expected five class folders, but each contains only a .gitkeep placeholder. The models directory is empty. The training notebook contains unexecuted cells and no output. Therefore, there is no evidence that an epoch completed, that a checkpoint was saved, or that a test set was evaluated. TensorFlow is not installed in the inspected environment, providing an additional reason why training could not be reproduced here.")
    add_heading(doc, "4.4 Quantitative Results", 2)
    add_para(doc, "No quantitative values can be reported responsibly. The table below is intentionally marked unavailable rather than populated with estimated or fabricated values. It should be completed only from the JSON files and classification reports produced by a recorded evaluation run.")
    add_table(doc, ["Metric", "Baseline CNN", "MobileNetV2", "Evidence required"], [
        ["Test accuracy", "Not available", "Not available", "*_metrics.json after evaluation"],
        ["Macro precision", "Not available", "Not available", "*_metrics.json after evaluation"],
        ["Macro recall", "Not available", "Not available", "*_metrics.json after evaluation"],
        ["Macro F1-score", "Not available", "Not available", "*_metrics.json after evaluation"],
        ["Average inference time", "Not available", "Not available", "*_metrics.json plus hardware details"],
        ["Per-class results", "Not available", "Not available", "Saved console report or CSV"],
        ["Confusion matrix", "Not available", "Not available", "*_confusion_matrix.png"],
    ])
    add_heading(doc, "4.5 Interpretation", 2)
    add_para(doc, "No claim can be made about which model performs better, whether the classifier reaches an acceptable threshold, or whether it generalises to Nigerian field conditions. Such claims require a fixed and documented test set. Once results are available, accuracy should be interpreted alongside macro F1-score and class-level recall; a high accuracy paired with poor organic-waste recall, for example, would not support a conclusion that the system works equally well for all classes. The confusion matrix should guide a qualitative examination of the images involved in common confusions.")
    add_heading(doc, "4.6 Required Reproducible Results Run", 2)
    add_bullets(doc, [
        "Populate each split with non-duplicate images and preserve a manifest containing file name, class, source and split.",
        "Run the two training scripts in an environment with the listed dependencies, retaining terminal logs and training-curve images.",
        "Evaluate each exact saved .keras checkpoint once on the untouched test split.",
        "Archive the two metrics JSON files, confusion matrices, per-class reports, package versions and hardware details.",
        "Replace every ‘Not available’ cell in Table 4.2 with the archived measured value and insert the actual confusion matrices as figures.",
        "Retest the deployed interface using supported images only after selecting the best evidence-supported model."
    ])
    add_heading(doc, "4.7 Chapter Summary", 2)
    add_para(doc, "The software architecture and training/evaluation pipeline have been implemented, but the experimental stage is incomplete in the supplied repository. Chapter Five therefore provides conclusions about the current development status and recommendations for completing a valid empirical study, rather than unsupported conclusions about model accuracy.")

    chapter_title(doc, 5, "Conclusion and Recommendations")
    add_heading(doc, "5.1 Conclusion", 2)
    add_para(doc, "This project provides the code structure for an intelligent waste-classification prototype: a dataset preparation utility, a baseline CNN, a MobileNetV2 transfer-learning alternative, an evaluation script, automated-test files and two web interfaces. The design supports a legitimate future comparison of predictive quality and inference speed. However, the supplied project does not contain image data, trained model files, training histories, evaluation reports or confusion matrices. Accordingly, the study cannot yet conclude that either model has been trained, that either achieves a particular accuracy, or that either is suitable for deployment. Reporting such values would be misleading.")
    add_heading(doc, "5.2 Recommendations", 2)
    add_bullets(doc, [
        "Complete and document the dataset before training, including class counts, data sources, permissions and a duplicate/leakage check across splits.",
        "Use the same class order and split for both models; never tune hyperparameters using the held-out test set.",
        "Publish accuracy together with macro precision, macro recall, macro F1-score, class-level support, confusion matrices, model file size, parameter count and measured inference time.",
        "Include a separately labelled local/field test subset if the intended claim concerns Nigerian conditions; report it separately from public-benchmark results.",
        "Record the device, CPU/GPU, TensorFlow version, batch size and timing procedure used for all runtime measures.",
        "Keep the interface’s confidence score clearly labelled as a model probability, not a guarantee of correctness, and retain human review for uncertain or safety-critical decisions.",
        "After a complete reproducible run, revise Chapter Four using the generated outputs and update this conclusion to address the measured evidence."
    ])
    add_heading(doc, "5.3 Contribution and Future Work", 2)
    add_para(doc, "The present contribution is a transparent and reproducible implementation plan rather than a completed accuracy study. Future work may add calibration assessment, an abstention threshold for low-confidence images, explainability visualisations, object detection for mixed scenes, and field testing with operational waste handlers. Each extension should be evaluated against a documented protocol rather than inferred from results on a curated image collection.")

    doc.add_page_break()
    h = doc.add_heading("REFERENCES", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    references = [
        "Chapman, P., Clinton, J., Kerber, R., Khabaza, T., Reinartz, T., Shearer, C., & Wirth, R. (2000). CRISP-DM 1.0: Step-by-step data mining guide.",
        "Kingma, D. P., & Ba, J. (2015). Adam: A method for stochastic optimization. International Conference on Learning Representations. https://arxiv.org/abs/1412.6980",
        "Pan, S. J., & Yang, Q. (2010). A survey on transfer learning. IEEE Transactions on Knowledge and Data Engineering, 22(10), 1345–1359. https://doi.org/10.1109/TKDE.2009.191",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
        "Sandler, M., Howard, A., Zhu, M., Zhmoginov, A., & Chen, L.-C. (2018). MobileNetV2: Inverted residuals and linear bottlenecks. Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, 4510–4520. https://doi.org/10.1109/CVPR.2018.00474",
        "Thung, G., & Yang, M. (2016). Classification of trash for recyclability status. Stanford CS229 Project Report. https://github.com/garythung/trashnet",
        "United Nations Environment Programme. (2018). Africa waste management outlook. UNEP.",
        "World Bank. (2025). What a Waste 3.0: Global snapshot of solid waste management toward circularity until 2050. World Bank. https://www.worldbank.org/en/publication/what-a-waste"
    ]
    for ref in references:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(6)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
