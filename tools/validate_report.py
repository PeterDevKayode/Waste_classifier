from pathlib import Path
import json
import re
from docx import Document

root = Path(__file__).resolve().parents[1]
doc_path = root / "Revised_Project_Report_Chapters_1_to_5.docx"
doc = Document(doc_path)
paragraphs = [p.text for p in doc.paragraphs]
tables = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
full = "\n".join(paragraphs + tables)
metrics = json.loads((root / "models/baseline_cnn_waste_metrics.json").read_text())

checks = {
    "docx_exists": doc_path.exists(),
    "docx_bytes": doc_path.stat().st_size,
    "paragraphs": len(paragraphs),
    "tables": len(doc.tables),
    "accuracy_match": f"{metrics['accuracy']:.4f}" in full,
    "macro_precision_match": f"{metrics['precision_macro']:.4f}" in full,
    "macro_recall_match": f"{metrics['recall_macro']:.4f}" in full,
    "macro_f1_match": f"{metrics['f1_macro']:.4f}" in full,
    "inference_match": f"{metrics['avg_inference_ms']:.2f}" in full,
    "citation_match": "10.24432/C5SS4G" in full and "archive.ics.uci.edu/dataset/908" in full,
    "nigeria_limitation": bool(re.search(r"\bNigeria\b|\bNigerian\b", full, re.I)),
    "no_placeholder": "Not available" not in full,
    "realwaste": "RealWaste" in full,
}
for key, value in checks.items():
    print(f"{key}: {value}")
assert checks["docx_exists"]
assert checks["docx_bytes"] > 10_000
assert checks["accuracy_match"] and checks["macro_precision_match"]
assert checks["macro_recall_match"] and checks["macro_f1_match"]
assert checks["inference_match"] and checks["citation_match"]
assert checks["nigeria_limitation"] and checks["no_placeholder"] and checks["realwaste"]
